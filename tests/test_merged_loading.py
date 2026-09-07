from datetime import date

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

import master_workbook_orchestrator_merged as orchestrator
from master_workbook_orchestrator_merged import Combination


def _control_workbook():
    workbook = Workbook()
    control = workbook.active
    control.title = "Global Control"

    for column, value in enumerate(("Entity", "Category", "SourceSheet", "Selected"), 1):
        control.cell(1, column, value)
    control.append(("JRL", "Bonds", "JRL", True))
    control.add_table(Table(displayName="tbl_entity_mapping", ref="A1:D2"))

    for column, value in enumerate(
        ("Stress", "Transition Tab", "Spread Tab", "Suffix", "OutputFolder", "Selected"), 1
    ):
        control.cell(4, column, value)
    control.append(("1in20", "Transition Matrix", "Spread Matrix", "1in20", "outputs", True))
    control.add_table(Table(displayName="tbl_stress_selection", ref="A4:F5"))

    for column, value in enumerate(("Setting", "Value"), 1):
        control.cell(7, column, value)
    control.append(("ValuationDate", date(2026, 6, 30)))
    control.append(("risk_factor_path", "risk factors"))
    control.append(("bool_inv_exp", True))
    control.add_table(Table(displayName="tbl_run_settings", ref="A7:B10"))

    for name in (
        "JRL",
        "Transition Matrix",
        "Spread Matrix",
        "I. Sensitivity - Template",
        "I. Control",
        *(f"Static{index}" for index in range(8)),
    ):
        worksheet = workbook.create_sheet(name)
        worksheet["A1"] = name
    workbook["I. Sensitivity - Template"].sheet_state = "hidden"
    for name, coordinate in zip(
        orchestrator.CONTROL_INPUT_NAMES,
        ("D10", "D13", "D16", "D19", "D22", "D25", "D28"),
    ):
        workbook.defined_names.add(
            DefinedName(name, attr_text=f"'I. Control'!${coordinate[0]}${coordinate[1:]}")
        )
    return workbook


def test_merged_run_uses_selective_package_loading(tmp_path, monkeypatch):
    master_path = tmp_path / "master.xlsx"
    _control_workbook().save(master_path)

    def fail_full_load(*args, **kwargs):
        raise AssertionError("full openpyxl loading should not be used by the merged runner")

    monkeypatch.setattr(orchestrator, "load_workbook", fail_full_load)
    monkeypatch.setattr(
        orchestrator,
        "_copy_output",
        lambda *args, **kwargs: args[4] / "output.xlsx",
    )

    manifests = orchestrator.run(master_path)

    assert len(manifests) == 1


def test_xml_serialisation_preserves_ignorable_namespace_declarations():
    parts = {
        "xl/workbook.xml": (
            b'<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            b'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            b'xmlns:x15="http://schemas.microsoft.com/office/spreadsheetml/2010/11/main" '
            b'mc:Ignorable="x15"><sheet /></workbook>'
        )
    }
    root = orchestrator._read_xml(parts, "xl/workbook.xml")
    orchestrator._write_xml(parts, "xl/workbook.xml", root)

    serialised = parts["xl/workbook.xml"].decode("utf-8")
    assert 'xmlns:x15="http://schemas.microsoft.com/office/spreadsheetml/2010/11/main"' in serialised
    assert 'mc:Ignorable="x15"' in serialised


def test_output_name_uses_compact_valuation_year_and_month():
    combination = Combination(
        "JRL", "Bonds", "JRL", "1in20", "Transition Matrix", "Spread Matrix", "1in20Comb"
    )

    assert orchestrator.output_name(combination, date(2026, 6, 30)) == (
        "202606_JRL_Bonds_1in20_1in20Comb.xlsx"
    )