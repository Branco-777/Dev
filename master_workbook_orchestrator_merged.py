"""Generate stress-specific workbooks through one package-level runner.

The runner treats the master workbook as an input package. It validates the
``Global Control`` sheet, expands the selected entity/category and stress
combinations, and creates one filtered workbook for each combination. Workbook
XML is edited directly so large data worksheets do not need to be loaded into
``openpyxl`` during validation or copied through a slower object model.

The public workflow is:

1. Call :func:`run` with the master workbook path.
2. Read and validate the control tables with :func:`parse_control`.
3. Copy and update one workbook package for each returned combination.
4. Write a manifest beside each distinct stress output folder.

The master workbook is never modified. Output files are written to temporary
files and moved into place only after the package has been fully created.
"""

from __future__ import annotations

import argparse
from copy import deepcopy
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta
from io import BytesIO
import itertools
import json
import os
from pathlib import Path
import posixpath
import re
import tempfile
import time
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries


CONTROL_SHEET = "Global Control"
ENTITY_TABLE = "tbl_entity_mapping"
STRESS_TABLE = "tbl_stress_selection"
SETTINGS_TABLE = "tbl_run_settings"
EXPECTED_STATIC_SHEETS = 8
DEFAULT_DATA_SHEET_NAME = "I. Bonds Data"
DEFAULT_SENSITIVITY_SHEET_NAME = "I. Sensitivity"
SENSITIVITY_TEMPLATE_SHEET = "I. Sensitivity - Template"
CONTROL_OUTPUT_SHEET = "I. Control"
CONTROL_INPUT_NAMES = (
    "valuation_date",
    "entity_name",
    "risk_factor_path",
    "output_folder_path",
    "prm_LPIgeneratorWB",
    "output_suffix",
    "bool_inv_exp",
)
SETTING_ALIASES = {
    "ValuationDate": {"valuationdate", "valuation date"},
    "risk_factor_path": {"risk_factor_path", "risk factor stresses path"},
    "bool_inv_exp": {"bool_inv_exp", "calculate investment expenses"},
}
INVALID_FILENAME_CHARS = re.compile(r"[<>:\"/\\|?*\x00-\x1f]")
INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")

MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
XML_NAMESPACES = {
    "": MAIN_NS,
    "r": REL_NS,
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "x14ac": "http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac",
    "x15": "http://schemas.microsoft.com/office/spreadsheetml/2010/11/main",
    "x15ac": "http://schemas.microsoft.com/office/spreadsheetml/2010/11/ac",
    "xr": "http://schemas.microsoft.com/office/spreadsheetml/2014/revision",
    "xr2": "http://schemas.microsoft.com/office/spreadsheetml/2015/revision2",
    "xr3": "http://schemas.microsoft.com/office/spreadsheetml/2016/revision3",
    "xr6": "http://schemas.microsoft.com/office/spreadsheetml/2016/revision6",
    "xr10": "http://schemas.microsoft.com/office/spreadsheetml/2016/revision10",
    "xcalcf": "http://schemas.microsoft.com/office/spreadsheetml/2018/calcfeatures",
    "xlwcv": "http://schemas.microsoft.com/office/spreadsheetml/2024/workbookCompatibilityVersion",
    "pr": PACKAGE_REL_NS,
}
NS = {"main": MAIN_NS, "rel": REL_NS, "package": PACKAGE_REL_NS}
R_ID = f"{{{REL_NS}}}id"

for _prefix, _namespace in XML_NAMESPACES.items():
    ET.register_namespace(_prefix, _namespace)


class OrchestratorError(ValueError):
    """Raised when the master workbook violates the control-sheet contract."""


@dataclass(frozen=True)
class Combination:
    """Describe one entity, category, and stress output combination.

    Attributes:
        entity: Entity name written to the generated ``I. Control`` sheet.
        category: Entity category used to identify the output.
        source_sheet: Master-workbook sheet containing the selected data.
        stress: Human-readable stress name.
        transition_tab: Worksheet containing the transition matrix source.
        spread_tab: Worksheet containing the spread matrix source.
        stress_suffix: Suffix written to the generated control values and name.
        output_folder: Resolved destination folder for the output workbook.
        risk_factor_path: Path passed through to the generated control sheet.
        bool_inv_exp: Investment-expense setting passed through unchanged.
        output_folder_value: Original folder value from the control table.
    """

    entity: str
    category: str
    source_sheet: str
    stress: str
    transition_tab: str
    spread_tab: str
    stress_suffix: str = ""
    output_folder: Path | None = None
    risk_factor_path: str = ""
    bool_inv_exp: Any = True
    output_folder_value: str | None = None


@dataclass
class RunResult:
    """Record the result of one attempted output workbook.

    Attributes:
        combination: JSON-serialisable combination details.
        output_path: Created workbook path, or ``None`` after a failure.
        status: ``"success"`` or ``"failed"``.
        error: Failure message when ``status`` is ``"failed"``.
    """

    combination: dict[str, Any]
    output_path: str | None
    status: str
    error: str | None = None


@dataclass(frozen=True)
class _XmlCell:
    """Represent the value needed from one cell in a package worksheet."""

    value: Any = None


@dataclass(frozen=True)
class _XmlTable:
    """Represent the range reference needed from one Excel table."""

    ref: str


class _XmlWorksheet:
    """Expose the small worksheet interface required by control validation."""

    def __init__(
        self,
        sheet_state: str,
        tables: dict[str, _XmlTable],
        values: dict[str, Any],
    ) -> None:
        """Initialise worksheet metadata and parsed cell values.

        Args:
            sheet_state: Excel worksheet state, such as ``"visible"`` or
                ``"hidden"``.
            tables: Tables indexed by their display name.
            values: Parsed cell values indexed by Excel coordinate.
        """
        self.sheet_state = sheet_state
        self.tables = tables
        self._values = values

    def cell(self, row: int, column: int) -> _XmlCell:
        """Return the value at a one-based row and column coordinate."""
        coordinate = f"{get_column_letter(column)}{row}"
        return _XmlCell(self._values.get(coordinate))


class _XmlDefinedName:
    """Represent a workbook-scoped or worksheet-scoped defined name."""

    def __init__(self, name: str, value: str, local_sheet_id: int | None) -> None:
        """Initialise a defined name from workbook XML attributes.

        Args:
            name: Defined-name identifier.
            value: Formula or cell reference stored in the workbook.
            local_sheet_id: Worksheet scope, or ``None`` for workbook scope.
        """
        self.name = name
        self.value = value
        self.localSheetId = local_sheet_id

    @property
    def destinations(self) -> list[tuple[str, str]]:
        """Return simple worksheet/cell destinations referenced by the name."""
        match = re.fullmatch(r"'((?:[^']|'')+)'!((?:\$?[A-Z]{1,3})\$?\d+)", self.value)
        if match is None:
            match = re.fullmatch(r"([A-Za-z0-9_. ]+)!((?:\$?[A-Z]{1,3})\$?\d+)", self.value)
        if match is None:
            return []
        return [(match.group(1).replace("''", "'"), match.group(2))]


class _XmlDefinedNames:
    """Provide name lookup for parsed workbook defined names."""

    def __init__(self, values: dict[str, _XmlDefinedName]) -> None:
        """Store defined names indexed by their workbook name."""
        self._values = values

    def get(self, name: str) -> _XmlDefinedName | None:
        """Return a defined name, or ``None`` when it is not present."""
        return self._values.get(name)


class _XmlWorkbook:
    """Minimal workbook representation used during package-level validation."""

    def __init__(self, worksheets: dict[str, _XmlWorksheet], defined_names: _XmlDefinedNames) -> None:
        """Initialise worksheet metadata and defined-name lookup."""
        self._worksheets = worksheets
        self.sheetnames = list(worksheets)
        self.defined_names = defined_names

    def __getitem__(self, name: str) -> _XmlWorksheet:
        """Return a worksheet by its exact Excel sheet name."""
        return self._worksheets[name]


def _date_style_ids(parts: dict[str, bytes]) -> set[int]:
    """Find style IDs that should be interpreted as dates or date-times.

    Args:
        parts: XLSX package parts indexed by package path.

    Returns:
        Style IDs using built-in or custom date-like number formats.
    """
    styles = _read_xml(parts, "xl/styles.xml")
    date_format_ids = {14, 15, 16, 17, 18, 19, 20, 21, 22, 27, 28, 29, 30, 31, 32, 33, 34, 45, 46, 47}
    custom_formats = {
        int(number_format.attrib["numFmtId"]): number_format.attrib.get("formatCode", "")
        for number_format in styles.findall("main:numFmts/main:numFmt", NS)
    }
    date_style_ids: set[int] = set()
    cell_formats = styles.find("main:cellXfs", NS)
    if cell_formats is None:
        return date_style_ids
    for style_id, cell_format in enumerate(cell_formats):
        number_format_id = int(cell_format.attrib.get("numFmtId", "0"))
        format_code = custom_formats.get(number_format_id, "")
        if number_format_id in date_format_ids or re.search(r"[dy]", format_code, re.IGNORECASE):
            date_style_ids.add(style_id)
    return date_style_ids


def _shared_strings(parts: dict[str, bytes]) -> list[str]:
    """Read the shared-string table from an XLSX package, if present.

    Args:
        parts: XLSX package parts indexed by package path.

    Returns:
        Shared strings in the order used by worksheet cell references.
    """
    if "xl/sharedStrings.xml" not in parts:
        return []
    root = _read_xml(parts, "xl/sharedStrings.xml")
    return ["".join(item.text or "" for item in si.iter(_tag("t"))) for si in root.findall(_tag("si"))]


def _xml_cell_values(
    root: ET.Element,
    shared_strings: list[str],
    date_style_ids: set[int],
) -> dict[str, Any]:
    """Parse worksheet cell values needed by the control-table reader.

    Args:
        root: Worksheet XML root element.
        shared_strings: Workbook shared-string values.
        date_style_ids: Style IDs that represent dates or date-times.

    Returns:
        Cell values indexed by their Excel coordinates.
    """
    values: dict[str, Any] = {}
    for cell in root.iter(_tag("c")):
        coordinate = cell.attrib.get("r")
        if coordinate is None:
            continue
        value_node = cell.find(_tag("v"))
        cell_type = cell.attrib.get("t")
        if cell_type == "inlineStr":
            value: Any = "".join(item.text or "" for item in cell.iter(_tag("t")))
        elif cell_type == "s" and value_node is not None:
            value = shared_strings[int(value_node.text or "0")]
        elif cell_type == "b" and value_node is not None:
            value = value_node.text == "1"
        elif value_node is None:
            value = None
        else:
            numeric_value = float(value_node.text or "0")
            if cell.attrib.get("s") and int(cell.attrib["s"]) in date_style_ids:
                days = int(numeric_value)
                fraction = numeric_value - days
                value = date(1899, 12, 30) + timedelta(days=days)
                if fraction:
                    value = datetime.combine(value, datetime.min.time()) + timedelta(days=fraction)
            elif numeric_value.is_integer():
                value = int(numeric_value)
            else:
                value = numeric_value
        values[coordinate] = value
    return values


def _load_package_workbook(parts: dict[str, bytes]) -> _XmlWorkbook:
    """Build a lightweight workbook model from XLSX package metadata.

    Only the control worksheet's cells and tables are parsed. Other worksheet
    data is intentionally left in the package for direct copying later.

    Args:
        parts: XLSX package parts indexed by package path.

    Returns:
        Workbook metadata suitable for control validation.

    Raises:
        OrchestratorError: If the workbook has no worksheet collection.
    """
    workbook_root = _read_xml(parts, "xl/workbook.xml")
    relationships = _relationship_parts(parts, "xl/workbook.xml")
    shared_strings = _shared_strings(parts)
    date_style_ids = _date_style_ids(parts) if "xl/styles.xml" in parts else set()
    worksheets: dict[str, _XmlWorksheet] = {}
    sheets = workbook_root.find("main:sheets", NS)
    if sheets is None:
        raise OrchestratorError("Workbook has no worksheet collection")
    for sheet in sheets:
        sheet_name = sheet.attrib["name"]
        sheet_part = relationships[sheet.attrib[R_ID]][1]
        if sheet_name == CONTROL_SHEET:
            sheet_root = _read_xml(parts, sheet_part)
            tables = {
                table_name: _XmlTable(table.attrib["ref"])
                for table_name, (table, _) in _table_by_name(parts, sheet_part).items()
            }
            values = _xml_cell_values(sheet_root, shared_strings, date_style_ids)
        else:
            tables = {}
            values = {}
        worksheets[sheet_name] = _XmlWorksheet(sheet.attrib.get("state", "visible"), tables, values)
    defined_name_values: dict[str, _XmlDefinedName] = {}
    defined_names = workbook_root.find("main:definedNames", NS)
    if defined_names is not None:
        for defined_name in defined_names:
            name = defined_name.attrib.get("name")
            if name is None:
                continue
            local_sheet_id = defined_name.attrib.get("localSheetId")
            defined_name_values[name] = _XmlDefinedName(
                name,
                defined_name.text or "",
                int(local_sheet_id) if local_sheet_id is not None else None,
            )
    return _XmlWorkbook(worksheets, _XmlDefinedNames(defined_name_values))


def _table_rows(workbook: Any, table_name: str) -> list[dict[str, Any]]:
    """Read one control-table range as a list of row dictionaries.

    Args:
        workbook: Workbook-like object exposing ``sheetnames`` and worksheet
            tables.
        table_name: Required Excel table name on ``Global Control``.

    Returns:
        One dictionary per data row, keyed by the table headings.

    Raises:
        OrchestratorError: If the control sheet, table, or headings are invalid.
    """
    if CONTROL_SHEET not in workbook.sheetnames:
        raise OrchestratorError(f"Missing required worksheet: {CONTROL_SHEET}")
    worksheet = workbook[CONTROL_SHEET]
    if table_name not in worksheet.tables:
        raise OrchestratorError(f"Missing required table: {table_name}")
    table = worksheet.tables[table_name]
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    headers = [worksheet.cell(min_row, column).value for column in range(min_col, max_col + 1)]
    if any(not isinstance(header, str) or not header.strip() for header in headers):
        raise OrchestratorError(f"Table {table_name} has a blank column heading")
    return [
        dict(zip(headers, (worksheet.cell(row, column).value for column in range(min_col, max_col + 1))))
        for row in range(min_row + 1, max_row + 1)
    ]


def _required_text(row: dict[str, Any], key: str, table_name: str) -> str:
    """Return a required, trimmed text field from a control-table row.

    Args:
        row: Control-table row values.
        key: Required field name.
        table_name: Table name used in validation errors.

    Returns:
        The trimmed field value.

    Raises:
        OrchestratorError: If the field is missing, blank, or not text.
    """
    value = row.get(key)
    if not isinstance(value, str) or not value.strip():
        raise OrchestratorError(f"{table_name} requires a non-blank {key}")
    return value.strip()


def _selected(value: Any) -> bool:
    """Interpret the supported truthy values used by control tables."""
    return value is True or str(value).strip().lower() in {"yes", "true", "1", "selected"}


def _setting(rows: list[dict[str, Any]], name: str) -> Any:
    """Return exactly one run-setting value, allowing configured aliases.

    Args:
        rows: Rows from ``tbl_run_settings``.
        name: Canonical setting name.

    Returns:
        The configured setting value.

    Raises:
        OrchestratorError: If zero or multiple matching settings exist.
    """
    accepted_names = {value.casefold() for value in SETTING_ALIASES.get(name, {name})}
    values = [
        row.get("Value")
        for row in rows
        if str(row.get("Setting", "")).strip().casefold() in accepted_names
    ]
    if len(values) != 1:
        raise OrchestratorError(f"Run settings must contain exactly one {name} value")
    return values[0]


def _entity_source_sheets(workbook: Any) -> set[str]:
    """Validate and collect source worksheets configured for entities.

    Args:
        workbook: Workbook-like object containing the control tables.

    Returns:
        Unique source-sheet names referenced by entity mappings.

    Raises:
        OrchestratorError: If a mapping is incomplete or its sheet is missing.
    """
    source_sheets = set()
    for row in _table_rows(workbook, ENTITY_TABLE):
        source_sheet = _required_text(row, "SourceSheet", ENTITY_TABLE)
        if source_sheet not in workbook.sheetnames:
            raise OrchestratorError(f"Missing SourceSheet tab: {source_sheet}")
        source_sheets.add(source_sheet)
    return source_sheets


def _parse_date(value: Any) -> date:
    """Convert a control-sheet valuation date to a ``date`` instance.

    Args:
        value: Excel date, date-time, or ISO date value.

    Returns:
        The normalised valuation date.

    Raises:
        OrchestratorError: If ``value`` is not a supported date value.
    """
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError as error:
            raise OrchestratorError(f"Valuation date is invalid: {value}") from error
    raise OrchestratorError(f"Valuation date is invalid: {value}")


def _control_named_cells(workbook: Any) -> dict[str, str]:
    """Validate required named inputs and return their target cell addresses.

    Args:
        workbook: Workbook-like object containing defined names.

    Returns:
        Mapping from required input name to one cell on ``I. Control``.

    Raises:
        OrchestratorError: If the control sheet, names, scope, or destinations
            do not match the output-workbook contract.
    """
    if CONTROL_OUTPUT_SHEET not in workbook.sheetnames:
        raise OrchestratorError(f"Missing required worksheet: {CONTROL_OUTPUT_SHEET}")
    named_cells: dict[str, str] = {}
    for name in CONTROL_INPUT_NAMES:
        defined_name = workbook.defined_names.get(name)
        if defined_name is None:
            raise OrchestratorError(f"Missing required named input: {name}")
        if defined_name.localSheetId is not None:
            raise OrchestratorError(f"Named input must be workbook-scoped: {name}")
        destinations = list(defined_name.destinations)
        if len(destinations) != 1 or destinations[0][0] != CONTROL_OUTPUT_SHEET:
            raise OrchestratorError(f"Named input must refer to I. Control: {name}")
        destination = destinations[0][1]
        if ":" in destination or "!" in destination:
            raise OrchestratorError(f"Named input must refer to one cell on I. Control: {name}")
        named_cells[name] = destination.replace("$", "")
    return named_cells


def parse_control(workbook: Any, base_folder: Path | None = None) -> tuple[date, list[Combination]]:
    """Validate control tables and expand them into output combinations.

    Entity mappings and selected stresses are combined as a Cartesian product.
    Stress output folders are resolved relative to ``base_folder`` when the
    control value is relative.

    Args:
        workbook: Workbook-like object containing the control sheet and tabs.
        base_folder: Folder used to resolve relative stress output folders.
            Defaults to the current working directory when omitted.

    Returns:
        A pair containing the valuation date and all selected combinations.

    Raises:
        OrchestratorError: If any required table value, worksheet, setting, or
            selection is invalid.
    """
    mappings = _table_rows(workbook, ENTITY_TABLE)
    stresses = _table_rows(workbook, STRESS_TABLE)
    settings = _table_rows(workbook, SETTINGS_TABLE)
    _entity_source_sheets(workbook)

    selected_mappings: list[tuple[str, str, str]] = []
    seen_mappings: set[tuple[str, str]] = set()
    for row in mappings:
        entity = _required_text(row, "Entity", ENTITY_TABLE)
        category = _required_text(row, "Category", ENTITY_TABLE)
        source_sheet = _required_text(row, "SourceSheet", ENTITY_TABLE)
        key = (entity, category)
        if key in seen_mappings:
            raise OrchestratorError(f"Duplicate entity/category mapping: {entity}/{category}")
        seen_mappings.add(key)
        if _selected(row.get("Selected")):
            selected_mappings.append((entity, category, source_sheet))

    selected_stresses: list[tuple[str, str, str, str, Path, str]] = []
    seen_stresses: set[str] = set()
    for row in stresses:
        stress = _required_text(row, "Stress", STRESS_TABLE)
        transition_tab = _required_text(row, "Transition Tab", STRESS_TABLE)
        spread_tab = _required_text(row, "Spread Tab", STRESS_TABLE)
        stress_suffix = _required_text(row, "Suffix", STRESS_TABLE)
        if stress in seen_stresses:
            raise OrchestratorError(f"Duplicate stress: {stress}")
        seen_stresses.add(stress)
        for tab_name, label in ((transition_tab, "Transition Tab"), (spread_tab, "Spread Tab")):
            if tab_name not in workbook.sheetnames:
                raise OrchestratorError(f"Missing {label}: {tab_name}")
        if _selected(row.get("Selected")):
            output_value = row.get("OutputFolder")
            if not isinstance(output_value, str) or not output_value.strip():
                raise OrchestratorError(
                    f"{STRESS_TABLE} requires a non-blank OutputFolder for selected stress: {stress}"
                )
            output_folder = Path(output_value.strip()).expanduser()
            if not output_folder.is_absolute():
                output_folder = (base_folder or Path.cwd()) / output_folder
            selected_stresses.append(
                (stress, transition_tab, spread_tab, stress_suffix, output_folder, output_value.strip())
            )

    if not selected_mappings:
        raise OrchestratorError("At least one entity/category mapping must be selected")
    if not selected_stresses:
        raise OrchestratorError("At least one stress must be selected")
    valuation_date = _parse_date(_setting(settings, "ValuationDate"))
    has_global_output_folder = any(
        str(row.get("Setting", "")).strip() == "OutputFolder"
        for table_name in workbook[CONTROL_SHEET].tables
        for row in _table_rows(workbook, table_name)
    )
    if has_global_output_folder:
        raise OrchestratorError("Run settings must not contain global OutputFolder")
    risk_factor_path = _setting(settings, "risk_factor_path")
    if not isinstance(risk_factor_path, str) or not risk_factor_path.strip():
        raise OrchestratorError("Run settings require a non-blank risk_factor_path value")
    bool_inv_exp = _setting(settings, "bool_inv_exp")
    combinations = [
        Combination(
            entity,
            category,
            source_sheet,
            stress,
            transition_tab,
            spread_tab,
            stress_suffix,
            output_folder,
            risk_factor_path.strip(),
            bool_inv_exp,
            output_folder_value,
        )
        for (entity, category, source_sheet), (
            stress,
            transition_tab,
            spread_tab,
            stress_suffix,
            output_folder,
            output_folder_value,
        ) in itertools.product(selected_mappings, selected_stresses)
    ]
    return valuation_date, combinations


def static_sheet_names(workbook: Any, excluded: set[str] | None = None) -> list[str]:
    """Derive the static worksheets retained in every generated workbook.

    The control sheet, output control sheet, configured source and matrix tabs,
    sensitivity sheets, and explicitly excluded names are removed. The
    remaining set must contain the expected number of static worksheets.

    Args:
        workbook: Workbook-like object containing sheet names and control tables.
        excluded: Additional sheet names to exclude for the current output.

    Returns:
        Static sheet names in their original workbook order.

    Raises:
        OrchestratorError: If the derived count is not the expected count.
    """
    stress_tabs = {
        tab_name
        for row in _table_rows(workbook, STRESS_TABLE)
        for tab_name in (row.get("Transition Tab"), row.get("Spread Tab"))
        if isinstance(tab_name, str) and tab_name.strip()
    }
    source_tabs = _entity_source_sheets(workbook)
    static = [
        name
        for name in workbook.sheetnames
        if name != CONTROL_SHEET
        and name not in (excluded or set())
        and name not in source_tabs
        and name != CONTROL_OUTPUT_SHEET
        and name not in stress_tabs
        and not (name.startswith("Sensitivities - ") or name.startswith("Sensitivities_"))
    ]
    if len(static) != EXPECTED_STATIC_SHEETS:
        raise OrchestratorError(f"Expected {EXPECTED_STATIC_SHEETS} static sheets, found {len(static)}")
    return static


def sanitise_filename(value: str) -> str:
    """Replace characters that are invalid in Windows filenames.

    Args:
        value: Candidate filename component.

    Returns:
        A safe, non-blank filename component.
    """
    result = INVALID_FILENAME_CHARS.sub("_", value).strip().rstrip(".")
    return result or "unnamed"


def output_name(combination: Combination, valuation_date: date) -> str:
    """Build the deterministic filename for one output combination.

    Args:
        combination: Entity, category, and stress suffix details for the output.
        valuation_date: Date included in the filename prefix.

    Returns:
        Sanitised workbook filename in the format
        ``YYYYMM_<entity>_<category>_<suffix>.xlsx``. The stress name is not
        included because the suffix identifies the selected stress output.
    """
    parts = [valuation_date.strftime("%Y%m"), combination.entity, combination.category]
    if combination.stress_suffix:
        parts.append(combination.stress_suffix)
    return "_".join(sanitise_filename(part) for part in parts) + ".xlsx"


def _validate_output_sheet_names(
    data_sheet_name: str,
    sensitivity_sheet_name: str,
    static_names: list[str],
) -> None:
    """Validate requested output names against Excel and retained-sheet rules.

    Args:
        data_sheet_name: Name assigned to the selected source sheet.
        sensitivity_sheet_name: Name assigned to the generated sensitivity sheet.
        static_names: Static worksheet names retained in the output.

    Raises:
        OrchestratorError: If a name is blank, too long, invalid, duplicated,
            or conflicts with a retained worksheet.
    """
    names = (data_sheet_name, sensitivity_sheet_name)
    if any(not name.strip() for name in names):
        raise OrchestratorError("Output sheet names must be non-blank")
    if any(len(name) > 31 for name in names):
        raise OrchestratorError("Output sheet names must be 31 characters or fewer")
    if any(INVALID_SHEET_CHARS.search(name) for name in names):
        raise OrchestratorError("Output sheet names contain invalid characters")
    if data_sheet_name.casefold() == sensitivity_sheet_name.casefold():
        raise OrchestratorError("Output sheet names must be unique")
    retained_names = {CONTROL_OUTPUT_SHEET, *static_names}
    for name in names:
        if name.casefold() in {retained.casefold() for retained in retained_names}:
            raise OrchestratorError(f"Output sheet name conflicts with an existing sheet: {name}")


def _tag(name: str) -> str:
    """Return an XML tag qualified with the spreadsheet namespace."""
    return f"{{{MAIN_NS}}}{name}"


def _local_name(tag: str) -> str:
    """Return the local name from a namespace-qualified XML tag."""
    return tag.rsplit("}", 1)[-1]


def _relationship_target(source_part: str, target: str) -> str:
    """Resolve a package relationship target relative to its source part."""
    return posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target)).lstrip("/")


def _read_xml(parts: dict[str, bytes], name: str) -> ET.Element:
    """Parse one XML package part and return its root element."""
    return ET.fromstring(parts[name])


def _write_xml(parts: dict[str, bytes], name: str, root: ET.Element) -> None:
    """Serialise an XML root back into its named package part.

    The namespace declarations used by Excel's markup-compatibility attributes
    are restored when ``ElementTree`` omits them during serialisation.
    """
    serialised = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    ignorable = re.search(rb"\bmc:Ignorable=\"([^\"]+)\"", serialised)
    if ignorable is not None:
        root_start = serialised.find(b"<", serialised.find(b"?>") + 2)
        root_end = serialised.find(b">", root_start)
        declarations = b"".join(
            f' xmlns:{prefix}="{XML_NAMESPACES[prefix]}"'.encode("ascii")
            for prefix in ignorable.group(1).decode("ascii").split()
            if prefix in XML_NAMESPACES
            and f"xmlns:{prefix}=".encode("ascii") not in serialised[root_start:root_end]
        )
        serialised = serialised[:root_end] + declarations + serialised[root_end:]
    parts[name] = serialised


def _relationship_parts(parts: dict[str, bytes], part: str) -> dict[str, tuple[str, str]]:
    """Read relationships for one package part.

    Args:
        parts: XLSX package parts indexed by package path.
        part: Package path whose ``.rels`` file should be read.

    Returns:
        Relationship ID mapped to relationship type and resolved target path.
        An empty dictionary is returned when no relationship file exists.
    """
    rels_name = f"{posixpath.dirname(part)}/_rels/{posixpath.basename(part)}.rels"
    if rels_name not in parts:
        return {}
    root = _read_xml(parts, rels_name)
    return {
        relationship.attrib["Id"]: (
            relationship.attrib["Type"],
            _relationship_target(part, relationship.attrib["Target"]),
        )
        for relationship in root.findall(f"{{{PACKAGE_REL_NS}}}Relationship")
    }


def _sheet_parts(parts: dict[str, bytes]) -> dict[str, str]:
    """Map workbook worksheet names to their package part paths."""
    workbook = _read_xml(parts, "xl/workbook.xml")
    relationships = _relationship_parts(parts, "xl/workbook.xml")
    sheets = workbook.find("main:sheets", NS)
    if sheets is None:
        return {}
    return {
        sheet.attrib["name"]: relationships[sheet.attrib[R_ID]][1]
        for sheet in sheets
    }


def _cell(root: ET.Element, coordinate: str) -> ET.Element:
    """Return or create a worksheet cell at an Excel coordinate.

    Args:
        root: Worksheet XML root element.
        coordinate: Cell coordinate such as ``"D10"``.

    Returns:
        The existing or newly created cell element.

    Raises:
        OrchestratorError: If ``coordinate`` does not contain a row number.
    """
    for cell in root.iter(_tag("c")):
        if cell.attrib.get("r") == coordinate:
            return cell
    sheet_data = root.find(_tag("sheetData"))
    if sheet_data is None:
        sheet_data = ET.SubElement(root, _tag("sheetData"))
    row_number_match = re.search(r"\d+", coordinate)
    if row_number_match is None:
        raise OrchestratorError(f"Invalid cell coordinate: {coordinate}")
    row_number = int(row_number_match.group())
    row = next((item for item in sheet_data if item.attrib.get("r") == str(row_number)), None)
    if row is None:
        row = ET.SubElement(sheet_data, _tag("row"), {"r": str(row_number)})
    return ET.SubElement(row, _tag("c"), {"r": coordinate})


def _set_cell_value(root: ET.Element, coordinate: str, value: Any) -> None:
    """Replace a cell's stored value while preserving unrelated cell metadata.

    Args:
        root: Worksheet XML root element.
        coordinate: Cell coordinate to update.
        value: Boolean, text, date, number, or ``None`` value to store.
    """
    cell = _cell(root, coordinate)
    for child in list(cell):
        if _local_name(child.tag) in {"f", "v", "is"}:
            cell.remove(child)
    cell.attrib.pop("t", None)
    if value is None:
        return
    if isinstance(value, bool):
        cell.attrib["t"] = "b"
        ET.SubElement(cell, _tag("v")).text = "1" if value else "0"
    elif isinstance(value, str):
        cell.attrib["t"] = "inlineStr"
        inline = ET.SubElement(cell, _tag("is"))
        ET.SubElement(inline, _tag("t")).text = value
    elif isinstance(value, date):
        serial = (value - date(1899, 12, 30)).days
        ET.SubElement(cell, _tag("v")).text = str(serial)
    else:
        ET.SubElement(cell, _tag("v")).text = str(value)


def _table_parts(parts: dict[str, bytes], sheet_part: str) -> list[tuple[ET.Element, str]]:
    """Return table XML roots and package paths linked to one worksheet."""
    worksheet = _read_xml(parts, sheet_part)
    relationships = _relationship_parts(parts, sheet_part)
    table_parts = worksheet.find("main:tableParts", NS)
    if table_parts is None:
        return []
    return [
        (table_root, target)
        for table_part in table_parts
        for target in [relationships[table_part.attrib[R_ID]][1]]
        for table_root in [_read_xml(parts, target)]
    ]


def _table_by_name(parts: dict[str, bytes], sheet_part: str) -> dict[str, tuple[ET.Element, str]]:
    """Index a worksheet's table XML by Excel table name."""
    return {
        table.attrib["name"]: (table, part)
        for table, part in _table_parts(parts, sheet_part)
    }


def _copy_value(source_root: ET.Element, target_root: ET.Element, source_coordinate: str, target_coordinate: str) -> None:
    """Copy a cell's value or formula from one worksheet XML tree to another."""
    source = _cell(source_root, source_coordinate)
    target = _cell(target_root, target_coordinate)
    for child in list(target):
        if _local_name(child.tag) in {"f", "v", "is"}:
            target.remove(child)
    target.attrib.pop("t", None)
    source_type = source.attrib.get("t")
    if source_type is not None:
        target.attrib["t"] = source_type
    for child in source:
        if _local_name(child.tag) in {"f", "v", "is"}:
            target.append(deepcopy(child))


def _populate_matrix(source_root: ET.Element, target_root: ET.Element, source_table: ET.Element, target_table: ET.Element) -> None:
    """Copy a matrix table cell-for-cell into a target table.

    Args:
        source_root: Worksheet containing the source matrix.
        target_root: Worksheet containing the target matrix tables.
        source_table: Source table element and range.
        target_table: Target table element and range.

    Raises:
        OrchestratorError: If source and target matrix dimensions differ.
    """
    source_min_col, source_min_row, source_max_col, source_max_row = range_boundaries(source_table.attrib["ref"])
    target_min_col, target_min_row, target_max_col, target_max_row = range_boundaries(target_table.attrib["ref"])
    source_shape = (source_max_col - source_min_col, source_max_row - source_min_row)
    target_shape = (target_max_col - target_min_col, target_max_row - target_min_row)
    if source_shape != target_shape:
        raise OrchestratorError(
            f"Matrix dimensions do not match: {source_table.attrib['ref']} and {target_table.attrib['ref']}"
        )
    for row_offset in range(source_max_row - source_min_row + 1):
        for column_offset in range(source_max_col - source_min_col + 1):
            source_coordinate = f"{get_column_letter(source_min_col + column_offset)}{source_min_row + row_offset}"
            target_coordinate = f"{get_column_letter(target_min_col + column_offset)}{target_min_row + row_offset}"
            _copy_value(source_root, target_root, source_coordinate, target_coordinate)


def _populate_sensitivity_package(parts: dict[str, bytes], sheet_parts: dict[str, str], combination: Combination) -> None:
    """Populate the hidden sensitivity template with selected stress matrices.

    Transition tables receive the matrix from ``combination.transition_tab``;
    spread and original-spread tables receive the matrix from
    ``combination.spread_tab``.

    Args:
        parts: Mutable XLSX package parts for one output.
        sheet_parts: Worksheet-name to package-part mapping.
        combination: Stress matrix selection for this output.

    Raises:
        OrchestratorError: If source tabs do not contain exactly one table, a
            table has an unknown purpose, or matrix dimensions differ.
    """
    target_sheet = sheet_parts[SENSITIVITY_TEMPLATE_SHEET]
    transition_sheet = sheet_parts[combination.transition_tab]
    spread_sheet = sheet_parts[combination.spread_tab]
    target_root = _read_xml(parts, target_sheet)
    transition_root = _read_xml(parts, transition_sheet)
    spread_root = _read_xml(parts, spread_sheet)
    sensitivity_tables = _table_by_name(parts, target_sheet)
    transition_tables = _table_by_name(parts, transition_sheet)
    spread_tables = _table_by_name(parts, spread_sheet)
    if len(transition_tables) != 1 or len(spread_tables) != 1:
        raise OrchestratorError("Each matrix source tab must contain exactly one table")
    for table, _ in sensitivity_tables.values():
        table_name = table.attrib["name"].lower()
        if "transition" in table_name:
            source_root = transition_root
            source = next(iter(transition_tables.values()))[0]
        elif "spread" in table_name:
            source_root = spread_root
            source = next(iter(spread_tables.values()))[0]
        else:
            raise OrchestratorError(f"Unrecognised sensitivity table: {table.attrib['name']}")
        _populate_matrix(source_root, target_root, source, table)
    _write_xml(parts, target_sheet, target_root)


def _rename_sheet_references(formula: str, renames: dict[str, str]) -> str:
    """Rewrite worksheet references in a formula according to ``renames``."""
    for old_name, new_name in renames.items():
        pattern = re.compile(rf"(?<![A-Za-z0-9_])(?:'{re.escape(old_name)}'|{re.escape(old_name)})!")
        formula = pattern.sub(f"'{new_name}'!", formula)
    return formula


def _has_sheet_reference(formula: str, sheet_name: str) -> bool:
    """Return whether a formula contains a reference to ``sheet_name``."""
    pattern = re.compile(rf"(?:'{re.escape(sheet_name)}'|{re.escape(sheet_name)})!")
    return bool(pattern.search(formula))


def _rewrite_worksheet_formulas(
    parts: dict[str, bytes],
    sheet_parts: dict[str, str],
    retained_names: set[str],
    renames: dict[str, str],
) -> None:
    """Update formulas on retained worksheets after sheet names change.

    Args:
        parts: Mutable XLSX package parts for one output.
        sheet_parts: Worksheet-name to package-part mapping.
        retained_names: Worksheets that remain in the generated workbook.
        renames: Old-to-new worksheet names used by the output.
    """
    for sheet_name in retained_names:
        part = sheet_parts.get(sheet_name)
        if part is None:
            continue
        root = _read_xml(parts, part)
        changed = False
        for formula in root.iter(_tag("f")):
            if formula.text:
                updated = _rename_sheet_references(formula.text, renames)
                if updated != formula.text:
                    formula.text = updated
                    changed = True
        if changed:
            _write_xml(parts, part, root)


def _update_workbook(
    parts: dict[str, bytes],
    combination: Combination,
    static_names: list[str],
    named_cells: dict[str, str],
    data_sheet_name: str,
    sensitivity_sheet_name: str,
) -> dict[str, str]:
    """Filter and rename worksheets in one output workbook package.

    The selected data sheet, control sheet, sensitivity template, and static
    sheets are retained. All other worksheets and worksheet relationships are
    removed. Formulas and workbook-scoped names are then updated to use the new
    output sheet names; names depending on removed worksheets are discarded.

    Args:
        parts: Mutable XLSX package parts for one output.
        combination: Entity and source-sheet selection for this output.
        static_names: Static worksheet names to retain.
        named_cells: Required control input names and their cell addresses.
        data_sheet_name: Output name for the selected source sheet.
        sensitivity_sheet_name: Output name for the sensitivity template.

    Returns:
        Worksheet-name to package-part mapping captured before renaming.

    Raises:
        OrchestratorError: If workbook XML is missing required worksheet data or
            a named control coordinate is invalid.
    """
    workbook = _read_xml(parts, "xl/workbook.xml")
    relationships = _read_xml(parts, "xl/_rels/workbook.xml.rels")
    sheet_parts = _sheet_parts(parts)
    retained = {combination.source_sheet, CONTROL_OUTPUT_SHEET, SENSITIVITY_TEMPLATE_SHEET, *static_names}
    sheet_nodes = workbook.find("main:sheets", NS)
    if sheet_nodes is None:
        raise OrchestratorError("Workbook has no worksheet collection")
    retained_relationship_ids: set[str] = set()
    for sheet in list(sheet_nodes):
        name = sheet.attrib["name"]
        if name not in retained:
            sheet_nodes.remove(sheet)
            continue
        retained_relationship_ids.add(sheet.attrib[R_ID])
        if name == combination.source_sheet:
            sheet.attrib["name"] = data_sheet_name
        elif name == SENSITIVITY_TEMPLATE_SHEET:
            sheet.attrib["name"] = sensitivity_sheet_name
            sheet.attrib.pop("state", None)
    for relationship in list(relationships):
        if (
            relationship.attrib.get("Type", "").endswith("/worksheet")
            and relationship.attrib.get("Id") not in retained_relationship_ids
        ):
            relationships.remove(relationship)

    renames = {
        combination.source_sheet: data_sheet_name,
        SENSITIVITY_TEMPLATE_SHEET: sensitivity_sheet_name,
    }
    _rewrite_worksheet_formulas(parts, sheet_parts, retained, renames)
    defined_names = workbook.find("main:definedNames", NS)
    if defined_names is not None:
        for defined_name in list(defined_names):
            name = defined_name.attrib.get("name")
            if name in CONTROL_INPUT_NAMES:
                coordinate = named_cells[name]
                match = re.fullmatch(r"([A-Z]+)(\d+)", coordinate.replace("$", ""))
                if match is None:
                    raise OrchestratorError(f"Invalid control coordinate for {name}: {coordinate}")
                defined_name.text = f"'{CONTROL_OUTPUT_SHEET}'!${match.group(1)}${match.group(2)}"
                continue
            formula = defined_name.text or ""
            if any(
                _has_sheet_reference(formula, sheet_name)
                for sheet_name in sheet_parts
                if sheet_name not in retained
            ):
                defined_names.remove(defined_name)
                continue
            defined_name.text = _rename_sheet_references(formula, renames)
    _write_xml(parts, "xl/workbook.xml", workbook)
    _write_xml(parts, "xl/_rels/workbook.xml.rels", relationships)
    return sheet_parts


def _control_values(combination: Combination, valuation_date: date) -> dict[str, Any]:
    """Build values written to the generated ``I. Control`` sheet.

    Args:
        combination: Selected entity, stress, and control settings.
        valuation_date: Validated valuation date for the output.

    Returns:
        Mapping from required named input to its output cell value.
    """
    return {
        "valuation_date": valuation_date,
        "entity_name": combination.entity,
        "risk_factor_path": combination.risk_factor_path,
        "output_folder_path": combination.output_folder_value or str(combination.output_folder),
        "prm_LPIgeneratorWB": None,
        "output_suffix": combination.stress_suffix,
        "bool_inv_exp": combination.bool_inv_exp,
    }


def _copy_output(
    master_bytes: bytes,
    workbook: Any,
    combination: Combination,
    valuation_date: date,
    output_folder: Path,
    data_sheet_name: str,
    sensitivity_sheet_name: str,
    named_cells: dict[str, str] | None = None,
    master_parts: dict[str, bytes] | None = None,
) -> Path:
    """Create one filtered output workbook from the master package.

    The package is copied in memory, updated for the requested combination, and
    written to a temporary file in the destination folder. ``os.replace`` then
    publishes the completed workbook atomically.

    Args:
        master_bytes: Original master workbook bytes.
        workbook: Lightweight validated workbook model.
        combination: Entity, category, and stress selection to generate.
        valuation_date: Date written to the output control sheet and filename.
        output_folder: Destination folder for the workbook.
        data_sheet_name: Output name for the selected source worksheet.
        sensitivity_sheet_name: Output name for the sensitivity worksheet.
        named_cells: Optional prevalidated control-name cell mapping.
        master_parts: Optional decompressed package parts shared across outputs.

    Returns:
        Path to the published output workbook.

    Raises:
        FileExistsError: If the deterministic output path already exists.
        OrchestratorError: If required workbook sheets or package structures are
            invalid.
    """
    named_cells = named_cells or _control_named_cells(workbook)
    if SENSITIVITY_TEMPLATE_SHEET not in workbook.sheetnames:
        raise OrchestratorError(f"Missing sensitivity template sheet: {SENSITIVITY_TEMPLATE_SHEET}")
    if workbook[SENSITIVITY_TEMPLATE_SHEET].sheet_state not in {"hidden", "veryHidden"}:
        raise OrchestratorError(
            f"Sensitivity template sheet must be hidden: {SENSITIVITY_TEMPLATE_SHEET}"
        )
    source_names = {combination.source_sheet, combination.transition_tab, combination.spread_tab}
    missing = [name for name in source_names if name not in workbook.sheetnames]
    if missing:
        raise OrchestratorError(f"Missing worksheets for combination: {', '.join(sorted(missing))}")
    static_names = static_sheet_names(workbook, source_names | {SENSITIVITY_TEMPLATE_SHEET})
    _validate_output_sheet_names(data_sheet_name, sensitivity_sheet_name, static_names)
    output_folder.mkdir(parents=True, exist_ok=True)
    destination = output_folder / output_name(combination, valuation_date)
    if destination.exists():
        raise FileExistsError(f"Output already exists: {destination}")

    if master_parts is None:
        with ZipFile(BytesIO(master_bytes)) as source_package:
            master_parts = {name: source_package.read(name) for name in source_package.namelist()}
    parts = dict(master_parts)  # shallow copy: per-output edits must not leak into other outputs
    sheet_parts = _update_workbook(
        parts,
        combination,
        static_names,
        named_cells,
        data_sheet_name,
        sensitivity_sheet_name,
    )
    control_root = _read_xml(parts, sheet_parts[CONTROL_OUTPUT_SHEET])
    for name, value in _control_values(combination, valuation_date).items():
        _set_cell_value(control_root, named_cells[name], value)
    _write_xml(parts, sheet_parts[CONTROL_OUTPUT_SHEET], control_root)
    _populate_sensitivity_package(parts, sheet_parts, combination)

    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb", suffix=".xlsx", prefix=f".{destination.stem}-", dir=output_folder, delete=False
        ) as temporary:
            temporary_path = Path(temporary.name)
            with ZipFile(temporary, "w", ZIP_DEFLATED) as package:
                for name, content in parts.items():
                    package.writestr(name, content)
        os.replace(temporary_path, destination)
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()
    return destination


def _decompress_parts(master_bytes: bytes) -> dict[str, bytes]:
    """Decompress every package part once for reuse across generated outputs.

    Args:
        master_bytes: Complete master XLSX package.

    Returns:
        Package parts indexed by their ZIP member path.
    """
    with ZipFile(BytesIO(master_bytes)) as source_package:
        return {name: source_package.read(name) for name in source_package.namelist()}


def run(
    master_path: Path,
    overwrite: bool = False,
    data_sheet_name: str = DEFAULT_DATA_SHEET_NAME,
    sensitivity_sheet_name: str = DEFAULT_SENSITIVITY_SHEET_NAME,
) -> list[Path]:
    """Generate all selected workbooks and write per-folder manifests.

    The run is divided into loading, validation, preflight, output generation,
    and manifest-writing stages. Validation and collision checks happen before
    any output is created. Individual output failures are recorded in the
    relevant manifest while other combinations continue to run.

    Args:
        master_path: Path to the read-only master workbook.
        overwrite: Whether existing deterministic output files may be replaced.
        data_sheet_name: Name for the selected source sheet in each output.
        sensitivity_sheet_name: Name for the populated sensitivity sheet.

    Returns:
        Manifest paths, one for each distinct stress output folder.

    Raises:
        FileNotFoundError: If ``master_path`` does not exist.
        FileExistsError: If an output exists and ``overwrite`` is false.
        OSError: If the master or output folders cannot be read or written.
        OrchestratorError: If control values or workbook structure are invalid.
    """
    started = time.perf_counter()
    print(f"[1/5] Loading master workbook package: {master_path}")
    master_bytes = master_path.read_bytes()
    master_parts = _decompress_parts(master_bytes)
    workbook = _load_package_workbook(master_parts)
    print("[2/5] Validating control sheet and building combinations")
    named_cells = _control_named_cells(workbook)
    valuation_date, combinations = parse_control(workbook, master_path.parent)
    print(f"      Validated {len(combinations)} output combinations")
    print("[3/5] Preflighting output destinations")
    destinations: list[Path] = []
    for combination in combinations:
        source_names = {combination.source_sheet, combination.transition_tab, combination.spread_tab}
        static_names = static_sheet_names(workbook, source_names | {SENSITIVITY_TEMPLATE_SHEET})
        _validate_output_sheet_names(data_sheet_name, sensitivity_sheet_name, static_names)
        if combination.output_folder is None:
            raise OrchestratorError("Combination has no output folder")
        destinations.append(combination.output_folder / output_name(combination, valuation_date))
    if not overwrite:
        collisions = [destination for destination in destinations if destination.exists()]
        if collisions:
            raise FileExistsError(f"Output already exists: {collisions[0]}")

    print("[4/5] Preparing output package")
    results: list[RunResult] = []
    for index, (combination, destination) in enumerate(zip(combinations, destinations), start=1):
        print(
            f"      Generating {index}/{len(combinations)}: "
            f"{combination.entity} / {combination.category} / {combination.stress}"
        )
        try:
            if combination.output_folder is None:
                raise OrchestratorError("Combination has no output folder")
            if overwrite:
                destination.unlink(missing_ok=True)
            created = _copy_output(
                master_bytes,
                workbook,
                combination,
                valuation_date,
                combination.output_folder,
                data_sheet_name,
                sensitivity_sheet_name,
                named_cells,
                master_parts=master_parts,
            )
            result_combination = asdict(combination)
            result_combination["output_folder"] = str(combination.output_folder)
            results.append(RunResult(result_combination, str(created), "success"))
        except (FileExistsError, OSError, OrchestratorError) as error:
            result_combination = asdict(combination)
            result_combination["output_folder"] = str(combination.output_folder) if combination.output_folder else None
            results.append(RunResult(result_combination, None, "failed", str(error)))

    timestamp = datetime.now()
    manifests: list[Path] = []
    folders = sorted({item.output_folder for item in combinations if item.output_folder}, key=str)
    print("[5/5] Writing run manifests")
    for folder in folders:
        folder.mkdir(parents=True, exist_ok=True)
        manifest = folder / f"run_manifest_{timestamp.strftime('%Y%m%dT%H%M%S')}.json"
        folder_results = [
            asdict(result)
            for result in results
            if result.combination["output_folder"] == str(folder)
        ]
        manifest.write_text(
            json.dumps(
                {
                    "master_path": str(master_path),
                    "timestamp": timestamp.isoformat(timespec="seconds"),
                    "valuation_date": valuation_date.isoformat(),
                    "output_folder": str(folder),
                    "results": folder_results,
                },
                indent=2,
            ),
            encoding="utf-8",
        )
        manifests.append(manifest)
        print(f"      Manifest written to {manifest}")
    elapsed = time.perf_counter() - started
    print(f"Completed consolidated workbook orchestration in {elapsed:.2f} seconds")
    return manifests


def main() -> int:
    """Parse command-line arguments, run orchestration, and return an exit code.

    Returns:
        ``0`` when every generated output succeeds, otherwise ``1``. Argument
        and workbook errors are reported through ``argparse``.
    """
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--master", required=True, type=Path)
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--data-sheet-name", default=DEFAULT_DATA_SHEET_NAME)
    parser.add_argument("--sensitivity-sheet-name", default=DEFAULT_SENSITIVITY_SHEET_NAME)
    args = parser.parse_args()
    try:
        manifests = run(args.master, args.overwrite, args.data_sheet_name, args.sensitivity_sheet_name)
    except (FileNotFoundError, OrchestratorError, FileExistsError, OSError) as error:
        parser.error(str(error))
    failed = False
    for manifest_path in manifests:
        print(f"Manifest written to {manifest_path}")
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        failed = failed or any(result["status"] != "success" for result in manifest["results"])
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())