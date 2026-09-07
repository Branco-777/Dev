"""Experimental package-level master workbook orchestrator.

This runner keeps the established openpyxl orchestrator untouched. It loads the
master once for control validation, then edits copied XLSX package parts directly.
"""

from __future__ import annotations

import argparse
from copy import deepcopy
from dataclasses import asdict
from datetime import date
from io import BytesIO
import json
from pathlib import Path
import posixpath
import re
import time
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

from master_workbook_orchestrator import (
    CONTROL_OUTPUT_SHEET,
    CONTROL_SHEET,
    CONTROL_INPUT_NAMES,
    DEFAULT_DATA_SHEET_NAME,
    DEFAULT_SENSITIVITY_SHEET_NAME,
    SENSITIVITY_TEMPLATE_SHEET,
    Combination,
    OrchestratorError,
    output_name,
    parse_control,
    static_sheet_names,
)

MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"main": MAIN_NS, "rel": REL_NS, "package": PACKAGE_REL_NS}
R_ID = f"{{{REL_NS}}}id"


def _tag(name: str) -> str:
    return f"{{{MAIN_NS}}}{name}"


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _relationship_target(source_part: str, target: str) -> str:
    return posixpath.normpath(
        posixpath.join(posixpath.dirname(source_part), target)
    ).lstrip("/")


def _read_xml(parts: dict[str, bytes], name: str) -> ET.Element:
    return ET.fromstring(parts[name])


def _write_xml(parts: dict[str, bytes], name: str, root: ET.Element) -> None:
    parts[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _relationship_parts(parts: dict[str, bytes], part: str) -> dict[str, tuple[str, str]]:
    rels_name = f"{posixpath.dirname(part)}/_rels/{posixpath.basename(part)}.rels"
    if rels_name not in parts:
        return {}
    root = _read_xml(parts, rels_name)
    return {
        relationship.attrib["Id"]: (
            relationship.attrib["Type"],
            _relationship_target(part, relationship.attrib["Target"]),
        )
        for relationship in root.findall(f"{{{PACKAGE_REL_NS}}}Relationship")
    }


def _sheet_parts(parts: dict[str, bytes]) -> dict[str, str]:
    workbook = _read_xml(parts, "xl/workbook.xml")
    relationships = _relationship_parts(parts, "xl/workbook.xml")
    result: dict[str, str] = {}
    sheets = workbook.find("main:sheets", NS)
    if sheets is None:
        return result
    for sheet in sheets:
        relationship = relationships[sheet.attrib[R_ID]]
        result[sheet.attrib["name"]] = relationship[1]
    return result


def _cell(root: ET.Element, coordinate: str) -> ET.Element:
    for cell in root.iter(_tag("c")):
        if cell.attrib.get("r") == coordinate:
            return cell
    sheet_data = root.find(_tag("sheetData"))
    if sheet_data is None:
        sheet_data = ET.SubElement(root, _tag("sheetData"))
    row_number = int(re.search(r"\d+", coordinate).group())
    row = next((item for item in sheet_data if item.attrib.get("r") == str(row_number)), None)
    if row is None:
        row = ET.SubElement(sheet_data, _tag("row"), {"r": str(row_number)})
    return ET.SubElement(row, _tag("c"), {"r": coordinate})


def _set_cell_value(root: ET.Element, coordinate: str, value: Any) -> None:
    cell = _cell(root, coordinate)
    for child in list(cell):
        if _local_name(child.tag) in {"f", "v", "is"}:
            cell.remove(child)
    cell.attrib.pop("t", None)
    if value is None:
        return
    if isinstance(value, bool):
        cell.attrib["t"] = "b"
        ET.SubElement(cell, _tag("v")).text = "1" if value else "0"
    elif isinstance(value, str):
        cell.attrib["t"] = "inlineStr"
        inline = ET.SubElement(cell, _tag("is"))
        ET.SubElement(inline, _tag("t")).text = value
    elif isinstance(value, date):
        serial = (value - date(1899, 12, 30)).days
        ET.SubElement(cell, _tag("v")).text = str(serial)
    else:
        ET.SubElement(cell, _tag("v")).text = str(value)


def _copy_value(source_root: ET.Element, target_root: ET.Element, source_coordinate: str, target_coordinate: str) -> None:
    source = _cell(source_root, source_coordinate)
    target = _cell(target_root, target_coordinate)
    for child in list(target):
        if _local_name(child.tag) in {"f", "v", "is"}:
            target.remove(child)
    target.attrib.pop("t", None)
    source_type = source.attrib.get("t")
    if source_type is not None:
        target.attrib["t"] = source_type
    for child in source:
        if _local_name(child.tag) in {"f", "v", "is"}:
            target.append(deepcopy(child))


def _table_parts(parts: dict[str, bytes], sheet_part: str) -> list[tuple[ET.Element, str]]:
    worksheet = _read_xml(parts, sheet_part)
    relationships = _relationship_parts(parts, sheet_part)
    result = []
    table_parts = worksheet.find("main:tableParts", NS)
    if table_parts is None:
        return result
    for table_part in table_parts:
        target = relationships[table_part.attrib[R_ID]][1]
        table_root = _read_xml(parts, target)
        result.append((table_root, target))
    return result


def _table_by_name(parts: dict[str, bytes], sheet_part: str) -> dict[str, tuple[ET.Element, str]]:
    return {
        table.attrib["name"]: (table, part)
        for table, part in _table_parts(parts, sheet_part)
    }


def _rename_sheet_reference(formula: str, old_name: str, new_name: str) -> str:
    return formula.replace(f"'{old_name}'!", f"'{new_name}'!").replace(
        f"{old_name}!", f"'{new_name}'!"
    )


def _populate_matrix(source_root: ET.Element, target_root: ET.Element, source_table: ET.Element, target_table: ET.Element) -> None:
    source_min_col, source_min_row, source_max_col, source_max_row = range_boundaries(source_table.attrib["ref"])
    target_min_col, target_min_row, target_max_col, target_max_row = range_boundaries(target_table.attrib["ref"])
    if (source_max_col - source_min_col, source_max_row - source_min_row) != (
        target_max_col - target_min_col,
        target_max_row - target_min_row,
    ):
        raise OrchestratorError(
            f"Matrix dimensions do not match: {source_table.attrib['ref']} and {target_table.attrib['ref']}"
        )
    for row_offset in range(source_max_row - source_min_row + 1):
        for column_offset in range(source_max_col - source_min_col + 1):
            source_coordinate = f"{get_column_letter(source_min_col + column_offset)}{source_min_row + row_offset}"
            target_coordinate = f"{get_column_letter(target_min_col + column_offset)}{target_min_row + row_offset}"
            _copy_value(source_root, target_root, source_coordinate, target_coordinate)
def _populate_sensitivity_package(parts: dict[str, bytes], sheet_parts: dict[str, str], combination: Combination) -> None:
    target_sheet = sheet_parts[SENSITIVITY_TEMPLATE_SHEET]
    transition_sheet = sheet_parts[combination.transition_tab]
    spread_sheet = sheet_parts[combination.spread_tab]
    target_root = _read_xml(parts, target_sheet)
    transition_root = _read_xml(parts, transition_sheet)
    spread_root = _read_xml(parts, spread_sheet)
    sensitivity_tables = _table_by_name(parts, target_sheet)
    transition_tables = _table_by_name(parts, transition_sheet)
    spread_tables = _table_by_name(parts, spread_sheet)
    if len(transition_tables) != 1 or len(spread_tables) != 1:
        raise OrchestratorError("Each matrix source tab must contain exactly one table")
    for table, _ in sensitivity_tables.values():
        table_name = table.attrib["name"].lower()
        if "transition" in table_name:
            source_tables = transition_tables
            source_root = transition_root
        elif "spread" in table_name:
            source_tables = spread_tables
            source_root = spread_root
        else:
            raise OrchestratorError(f"Unrecognised sensitivity table: {table.attrib['name']}")
        source = next(iter(source_tables.values()))[0]
        _populate_matrix(source_root, target_root, source, table)
    _write_xml(parts, target_sheet, target_root)


def _update_workbook(parts: dict[str, bytes], combination: Combination, static_names: list[str], named_cells: dict[str, str], data_sheet_name: str, sensitivity_sheet_name: str) -> None:
    workbook = _read_xml(parts, "xl/workbook.xml")
    relationships = _read_xml(parts, "xl/_rels/workbook.xml.rels")
    sheet_parts = _sheet_parts(parts)
    retained = {combination.source_sheet, CONTROL_OUTPUT_SHEET, SENSITIVITY_TEMPLATE_SHEET, *static_names}
    sheet_nodes = workbook.find("main:sheets", NS)
    if sheet_nodes is None:
        raise OrchestratorError("Workbook has no worksheet collection")
    retained_relationship_ids: set[str] = set()
    for sheet in list(sheet_nodes):
        name = sheet.attrib["name"]
        if name not in retained:
            sheet_nodes.remove(sheet)
        else:
            retained_relationship_ids.add(sheet.attrib[R_ID])
            if name == combination.source_sheet:
                sheet.attrib["name"] = data_sheet_name
            elif name == SENSITIVITY_TEMPLATE_SHEET:
                sheet.attrib["name"] = sensitivity_sheet_name
                sheet.attrib.pop("state", None)
    worksheet_relationship_ids = {
        sheet.attrib[R_ID]
        for sheet in list(sheet_nodes)
        if sheet.attrib.get(R_ID) in retained_relationship_ids
    }
    for relationship in list(relationships):
        if relationship.attrib.get("Type", "").endswith("/worksheet") and relationship.attrib.get("Id") not in worksheet_relationship_ids:
            relationships.remove(relationship)
    defined_names = workbook.find("main:definedNames", NS)
    if defined_names is not None:
        for defined_name in list(defined_names):
            name = defined_name.attrib.get("name")
            if name in CONTROL_INPUT_NAMES:
                coordinate = named_cells[name]
                match = re.fullmatch(r"([A-Z]+)(\d+)", coordinate.replace("$", ""))
                if match is None:
                    raise OrchestratorError(f"Invalid control coordinate for {name}: {coordinate}")
                defined_name.text = f"'{CONTROL_OUTPUT_SHEET}'!${match.group(1)}${match.group(2)}"
                continue
            formula = defined_name.text or ""
            referenced_sheet = next(
                (
                    sheet_name
                    for sheet_name in sheet_parts
                    if f"'{sheet_name}'!" in formula or f"{sheet_name}!" in formula
                ),
                None,
            )
            if referenced_sheet is not None and referenced_sheet not in retained:
                defined_names.remove(defined_name)
                continue
            if referenced_sheet == combination.source_sheet:
                defined_name.text = _rename_sheet_reference(formula, referenced_sheet, data_sheet_name)
            elif referenced_sheet == SENSITIVITY_TEMPLATE_SHEET:
                defined_name.text = _rename_sheet_reference(formula, referenced_sheet, sensitivity_sheet_name)
    _write_xml(parts, "xl/workbook.xml", workbook)
    _write_xml(parts, "xl/_rels/workbook.xml.rels", relationships)


def _copy_output(master_bytes: bytes, workbook: Any, combination: Combination, valuation_date: date, output_folder: Path, data_sheet_name: str, sensitivity_sheet_name: str) -> Path:
    if not data_sheet_name or not sensitivity_sheet_name:
        raise OrchestratorError("Output sheet names must be non-blank")
    if len(data_sheet_name) > 31 or len(sensitivity_sheet_name) > 31:
        raise OrchestratorError("Output sheet names must be 31 characters or fewer")
    if data_sheet_name == sensitivity_sheet_name:
        raise OrchestratorError("Output sheet names must be unique")
    named_cells = workbook.defined_names
    named_coordinates = {
        name: list(named_cells[name].destinations)[0][1].replace("$", "")
        for name in CONTROL_INPUT_NAMES
    }
    static_names = static_sheet_names(workbook, {combination.source_sheet, combination.transition_tab, combination.spread_tab, SENSITIVITY_TEMPLATE_SHEET})
    parts: dict[str, bytes]
    with ZipFile(BytesIO(master_bytes)) as source_package:
        parts = {name: source_package.read(name) for name in source_package.namelist()}
    sheet_parts = _sheet_parts(parts)
    _update_workbook(parts, combination, static_names, named_coordinates, data_sheet_name, sensitivity_sheet_name)
    control_root = _read_xml(parts, sheet_parts[CONTROL_OUTPUT_SHEET])
    values = {
        "valuation_date": valuation_date,
        "entity_name": combination.entity,
        "risk_factor_path": combination.risk_factor_path,
        "output_folder_path": combination.output_folder_value or str(combination.output_folder),
        "prm_LPIgeneratorWB": None,
        "output_suffix": combination.stress_suffix,
        "bool_inv_exp": combination.bool_inv_exp,
    }
    for name, value in values.items():
        _set_cell_value(control_root, named_coordinates[name], value)
    _write_xml(parts, sheet_parts[CONTROL_OUTPUT_SHEET], control_root)
    _populate_sensitivity_package(parts, sheet_parts, combination)
    output_folder.mkdir(parents=True, exist_ok=True)
    destination = output_folder / output_name(combination, valuation_date)
    if destination.exists():
        raise FileExistsError(f"Output already exists: {destination}")
    with ZipFile(destination, "w", ZIP_DEFLATED) as package:
        for name, content in parts.items():
            package.writestr(name, content)
    return destination


def run(master_path: Path, overwrite: bool = False, data_sheet_name: str = DEFAULT_DATA_SHEET_NAME, sensitivity_sheet_name: str = DEFAULT_SENSITIVITY_SHEET_NAME) -> list[Path]:
    started = time.perf_counter()
    print(f"[1/4] Loading and validating master workbook: {master_path}")
    workbook = load_workbook(master_path, data_only=False)
    valuation_date, combinations = parse_control(workbook, master_path.parent)
    print(f"[2/4] Validated {len(combinations)} output combinations")
    master_buffer = BytesIO()
    workbook.save(master_buffer)
    master_bytes = master_buffer.getvalue()
    print("[3/4] Serialised master workbook package")
    results = []
    for index, combination in enumerate(combinations, start=1):
        if combination.output_folder is None:
            raise OrchestratorError("Combination has no output folder")
        print(
            f"      Generating {index}/{len(combinations)}: "
            f"{combination.entity} / {combination.category} / {combination.stress}"
        )
        destination = combination.output_folder / output_name(combination, valuation_date)
        if overwrite:
            destination.unlink(missing_ok=True)
        try:
            created = _copy_output(master_bytes, workbook, combination, valuation_date, combination.output_folder, data_sheet_name, sensitivity_sheet_name)
            result_combination = asdict(combination)
            result_combination["output_folder"] = str(combination.output_folder)
            results.append({"combination": result_combination, "output_path": str(created), "status": "success", "error": None})
        except (FileExistsError, OSError, OrchestratorError) as error:
            result_combination = asdict(combination)
            result_combination["output_folder"] = str(combination.output_folder)
            results.append({"combination": result_combination, "output_path": None, "status": "failed", "error": str(error)})
    manifests = []
    timestamp = __import__("datetime").datetime.now()
    print("[4/4] Writing run manifests")
    for folder in sorted({item.output_folder for item in combinations if item.output_folder}, key=str):
        folder.mkdir(parents=True, exist_ok=True)
        manifest = folder / f"run_manifest_{timestamp.strftime('%Y%m%dT%H%M%S')}.json"
        folder_results = [item for item in results if item["combination"]["output_folder"] == str(folder)]
        manifest.write_text(json.dumps({"master_path": str(master_path), "timestamp": timestamp.isoformat(timespec="seconds"), "valuation_date": valuation_date.isoformat(), "output_folder": str(folder), "results": folder_results}, indent=2), encoding="utf-8")
        print(f"      Manifest written to {manifest}")
        manifests.append(manifest)
    elapsed = time.perf_counter() - started
    print(f"Completed package-level orchestration in {elapsed:.2f} seconds")
    return manifests


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--master", required=True, type=Path)
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--data-sheet-name", default=DEFAULT_DATA_SHEET_NAME)
    parser.add_argument("--sensitivity-sheet-name", default=DEFAULT_SENSITIVITY_SHEET_NAME)
    args = parser.parse_args()
    manifests = run(args.master, args.overwrite, args.data_sheet_name, args.sensitivity_sheet_name)
    failed = False
    for path in manifests:
        print(f"Manifest written to {path}")
        manifest = json.loads(path.read_text(encoding="utf-8"))
        failed = failed or any(result["status"] != "success" for result in manifest["results"])
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
